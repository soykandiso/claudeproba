"""The document template engine (roadmap P5 s56).

The acceptance (brief §8): "financial and legal fields are marked as requiring the
applicant's real data, never invented". Templates are built here with python-docx, the way
Word writes them, runs split mid-placeholder included.
"""

import datetime as dt
import io
import re

import docx
import pytest

from app import create_app
from app.config import load_settings
from app.documents import engine


def template(
    *paragraphs: str, header: str | None = None, cell: str | None = None, split=None
) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if split:  # one placeholder across two runs, as Word does after an edit
        p = document.add_paragraph()
        p.add_run(split[0])
        p.add_run(split[1])
    if cell:
        document.add_table(rows=1, cols=1).cell(0, 0).text = cell
    if header:
        document.sections[0].header.is_linked_to_previous = False
        document.sections[0].header.paragraphs[0].text = header
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


FORM = template(
    "Барател: {{r applicant_name }}, {{r applicant_municipality }}.",
    "Повик: {{r call_title }}, рок {{r call_deadline }}.",
    "ЕДБ: {{r applicant_edb }}",
    "Барани средства: {{r requested_amount }}",
    "Одговорно лице: {{r responsible_person }}",
    "{{r signature_and_stamp }}",
)
VALUES = {
    "applicant_name": "Пекара Проба ДООЕЛ",
    "applicant_municipality": "Битола",
    "call_title": "Јавен повик за поддршка на микро претпријатија",
    "call_deadline": dt.date(2026, 11, 30),
}


def test_what_the_platform_knows_is_filled_and_formatted():
    done = engine.render(FORM, VALUES)
    text = engine.text_of(done.docx)
    assert "Барател: Пекара Проба ДООЕЛ, Битола." in text
    assert "рок 30.11.2026." in text
    assert done.filled == [
        "applicant_name",
        "applicant_municipality",
        "call_title",
        "call_deadline",
    ]


def test_money_people_and_signatures_are_never_filled_and_say_so():
    done = engine.render(FORM, VALUES)
    text = engine.text_of(done.docx)
    assert "Барани средства: [Пополнува барателот: Барани средства]" in text
    assert "Одговорно лице: [Пополнува барателот: Одговорно лице (име и презиме)]" in text
    assert "[Пополнува барателот: Потпис на одговорното лице и печат]" in text
    applicant = [u.key for u in done.unfilled if u.reason == "applicant"]
    assert applicant == ["requested_amount", "responsible_person", "signature_and_stamp"]


@pytest.mark.parametrize("field", ["requested_amount", "responsible_person", "signature_and_stamp"])
def test_handing_a_value_to_an_applicant_field_is_refused_as_an_invention(field):
    with pytest.raises(engine.InventedValue, match=field):
        engine.render(FORM, {**VALUES, field: "1.000.000"})


def test_an_empty_value_for_an_applicant_field_is_not_an_invention():
    assert engine.render(FORM, {**VALUES, "requested_amount": "", "responsible_person": None})


def test_a_missing_fact_is_marked_and_listed_not_left_blank():
    done = engine.render(FORM, VALUES)  # no ЕДБ given
    assert "ЕДБ: [Недостасува: ЕДБ на барателот]" in engine.text_of(done.docx)
    assert engine.Unfilled("applicant_edb", "ЕДБ на барателот", "missing") in done.unfilled


def test_every_marker_is_highlighted_and_bold():
    done = engine.render(FORM, VALUES)
    xml = docx.Document(io.BytesIO(done.docx)).element.xml
    runs = re.findall(r"<w:r>(.*?)</w:r>", xml, re.S)
    marked = [r for r in runs if "[Пополнува барателот" in r or "[Недостасува" in r]
    assert len(marked) == 4
    assert all(f'w:fill="{engine.MARK}"' in r and "<w:b/>" in r for r in marked)


def test_every_money_field_is_the_applicants():
    """Brief §8: financial figures are never auto-filled. No money field may have a source."""
    money = {k: f for k, f in engine.registry().items() if f.type == "money"}
    assert money and all(f.source == "applicant" for f in money.values())


@pytest.mark.parametrize(
    "bad, said",
    [
        ("{{r applicant_nickname }}", "not a field"),
        ("{{ applicant_name }}", "write it {{r applicant_name }}"),
        ("{% if x %}а{% endif %}", "no logic"),
    ],
)
def test_a_template_that_cannot_be_filled_safely_is_refused(bad, said):
    bad = template(bad)
    assert any(said in p for p in engine.check(bad))
    with pytest.raises(engine.TemplateRefused):
        engine.render(bad, {})


def test_a_placeholder_split_across_runs_is_still_found_and_filled():
    split = template(split=("Назив: {{r applicant_", "name }}"))
    assert engine.placeholders(split) == [("applicant_name", True)]
    assert "Назив: Пекара" in engine.text_of(
        engine.render(split, {"applicant_name": "Пекара"}).docx
    )


def test_tables_and_headers_are_filled_too():
    form = template(
        "x",
        cell="ЕМБС: {{r applicant_embs }}",
        header="{{r applicant_name }}, {{r document_date }}",
    )
    done = engine.render(form, {"applicant_name": "Проба", "document_date": dt.date(2026, 10, 15)})
    text = engine.text_of(done.docx)
    assert "Проба, 15.10.2026" in text and "ЕМБС: [Недостасува: ЕМБС на барателот]" in text


def test_the_applicants_words_are_printed_as_given():
    done = engine.render(template("{{r applicant_name }}"), {"applicant_name": "А & Б <ДОО>"})
    assert "А & Б <ДОО>" in engine.text_of(done.docx)


def test_values_for_fields_that_do_not_exist_are_refused():
    with pytest.raises(engine.TemplateRefused, match="do not exist"):
        engine.render(FORM, {**VALUES, "turnover": 1})


def test_the_check_command_fails_on_a_bad_template(tmp_path):
    good, bad = tmp_path / "good.docx", tmp_path / "bad.docx"
    good.write_bytes(FORM)
    bad.write_bytes(template("{{ call_title }}"))
    for form in (good, bad):
        form.with_suffix(".yaml").write_text("title_mk: Образец\n", encoding="utf-8")
    runner = create_app(load_settings(env="testing")).test_cli_runner()
    result = runner.invoke(args=["documents", "check", str(good), str(bad)])
    assert result.exit_code == 1 and "1 of 2 templates can be filled" in result.output
