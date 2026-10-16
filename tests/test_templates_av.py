"""Template set 1, the Employment Agency's forms for employers (roadmap P5 s58).

The templates are АВ's own Word forms (https://av.gov.mk/obrasci.nspx, 15.10.2026) with
fields placed into their blank answer cells by `ops/dev/build_av_templates.py`.
"""

import datetime as dt
import io
import re

import docx
import pytest

from app.documents import TEMPLATES, checklist, derive, engine

AV = sorted((TEMPLATES / "av").rglob("*.docx"))
IDENTITY = {
    "applicant_name": "Пример Пекара ДООЕЛ Битола",
    "applicant_embs": "7654321",
    "applicant_edb": "МК4002020000000",
    "applicant_address": "ул. Пример бр. 1, 7000 Битола",
    "applicant_municipality": "Битола",
    "applicant_activity": "10.71 Производство на леб",
}


def _fill(path):
    used = {k for k, _ in engine.placeholders(path.read_bytes())}
    return engine.render(path.read_bytes(), {k: v for k, v in IDENTITY.items() if k in used})


def test_the_three_employer_forms_are_there():
    assert [p.parent.name for p in AV] == ["oglas-praktikant", "ovlastuvanje-praktikant", "ppr"]


@pytest.mark.parametrize("path", AV, ids=lambda p: p.parent.name)
def test_each_template_can_be_filled_and_says_where_its_form_came_from(path):
    assert engine.check(path.read_bytes()) == []
    manifest = checklist.manifest_for(path)
    assert manifest.source_url.startswith("https://av.gov.mk/content/Обрасци/")
    assert re.fullmatch(r"[0-9a-f]{64}", manifest.source_sha256)
    assert manifest.retrieved == dt.date(2026, 10, 15) and manifest.fill_by_hand_mk


@pytest.mark.parametrize("path", AV, ids=lambda p: p.parent.name)
def test_the_employers_identity_is_filled_and_nothing_of_the_applicants_is(path):
    done = _fill(path)
    assert {"applicant_name", "applicant_embs"} <= set(done.filled)
    applicant = {u.key for u in done.unfilled if u.reason == "applicant"}
    money = {k for k, f in engine.registry().items() if f.type == "money"}
    used = {k for k, _ in engine.placeholders(path.read_bytes())}
    assert used & money <= applicant  # the salary and the trainee's allowance are never filled
    text = engine.text_of(done.docx)
    assert "Пример Пекара ДООЕЛ Битола" in text and "7654321" in text


def test_the_agencys_own_part_of_the_form_is_untouched():
    ppr = TEMPLATES / "av" / "ppr" / "2024-04-11.docx"
    text = engine.text_of(_fill(ppr).docx)
    assert "Пополнува Агенцијата за вработување" in text
    assert "1. новоотворено" in text  # the form's options stay the form's


def test_the_signature_lines_carry_their_marker():
    for slug in ("ppr", "ovlastuvanje-praktikant"):
        path = next(p for p in AV if p.parent.name == slug)
        assert "[Пополнува барателот: Потпис на одговорното лице и печат]" in engine.text_of(
            _fill(path).docx
        )


def test_the_package_checklist_says_the_work_club_once():
    documents = [(checklist.manifest_for(p), _fill(p)) for p in AV]
    notes = checklist.package_notes([m for m, _ in documents])
    printed = checklist.to_docx(
        checklist.build(documents), applicant="а", call="б", on=dt.date(2026, 10, 16), notes=notes
    )
    assert engine.text_of(printed).count("работниот клуб") == 1


# ------------------------------------------------------------------ deriving a template


def _form(*rows) -> docx.document.Document:
    document = docx.Document()
    table = document.add_table(rows=0, cols=2)
    for label, answer in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, answer
    return document


def test_a_field_goes_only_into_an_empty_answer_cell():
    document = _form(("Назив на работодавачот / Emri", ""), ("Начин / Forma", "1. Да 2. Не"))
    derive.place(document, "Назив на работодавачот", "applicant_name")
    assert document.tables[0].rows[0].cells[1].text == "{{r applicant_name }}"
    with pytest.raises(derive.FormChanged, match="no empty answer cell"):
        derive.place(document, "Начин", "contact_phone")
    assert document.tables[0].rows[1].cells[1].text == "1. Да 2. Не"


def test_a_label_the_form_no_longer_has_is_an_error_not_a_skip():
    with pytest.raises(derive.FormChanged, match="no such row"):
        derive.place(_form(("Назив", "")), "Телефон", "contact_phone")


def test_appending_keeps_the_label_and_adds_the_field_after_it():
    document = _form(("Седиште и адреса / Selia", "x"))
    derive.append(document, "Седиште и адреса", "applicant_address")
    cell = document.tables[0].rows[0].cells[0]
    assert cell.paragraphs[0].text == "Седиште и адреса / Selia"
    assert cell.paragraphs[-1].text == "{{r applicant_address }}"
    out = io.BytesIO()
    document.save(out)
    assert engine.check(out.getvalue()) == []
