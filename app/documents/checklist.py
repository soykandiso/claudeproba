"""What the applicant must still provide (roadmap P5 s57, brief §8).

Every package ships a checklist. It is made from two things, never written by hand:

1. **every field the engine left unfilled** (`engine.Rendered.unfilled`), across all the
   package's documents: one line per field, naming the documents it appears in. A field only
   the applicant can supply and a fact the platform does not have are both listed, in groups
   an applicant works through: data, figures, declarations, signatures;
2. **every attachment the templates' manifests name** (`<version>.yaml` beside each
   `.docx`), from the catalogue `config/document_attachments.yaml`, with how to obtain it and
   the call's own condition (`note_mk`) when it has one. Listed once, however many documents
   ask for it.

The checklist is stored on the package (`document_package.checklist`, `as_rows`) and printed
as the package's last document (`to_docx`). Its words are linted like a report's.
"""

import datetime as dt
import io
from dataclasses import dataclass, field
from functools import cache
from pathlib import Path
from typing import Literal

import docx
import yaml
from docx.shared import Pt
from pydantic import BaseModel, ConfigDict

from app.documents import engine
from app.reports.lint import find_banned
from app.web.format import mkdate

CATALOGUE = Path(__file__).resolve().parents[2] / "config" / "document_attachments.yaml"

Group = Literal["data", "money", "declaration", "signature", "attachment"]
HEADINGS: dict[str, str] = {
    "data": "Податоци што недостасуваат во документите",
    "money": "Финансиски податоци",
    "declaration": "Изјави",
    "signature": "Потписи и печат",
    "attachment": "Прилози",
}
BOX = "☐"


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)


class Attachment(_Strict):
    label_mk: str
    how_mk: str


class Asked(_Strict):
    """One attachment a template's call asks for, with the call's condition if any."""

    key: str
    note_mk: str | None = None


class Manifest(_Strict):
    """`<version>.yaml` beside a template: what the document is and what it needs attached.

    A template made from an institution's own form says where the form came from
    (`source_url`, `source_sha256`, `retrieved`), so a newer form is noticed and the template
    made again (`ops/dev/build_av_templates.py`). `fill_by_hand_mk` is one checklist line for the
    blank cells the form leaves to the applicant (the vacancy, the trainee's tasks).
    """

    title_mk: str
    attachments: list[Asked] = []
    source_url: str | None = None
    source_sha256: str | None = None
    retrieved: dt.date | None = None
    fill_by_hand_mk: str | None = None
    # Said once at the top of the checklist, however many documents carry it.
    package_note_mk: str | None = None


class ChecklistRefused(ValueError):
    """The checklist cannot be made as it stands; nothing was produced."""


@cache
def catalogue(path: Path = CATALOGUE) -> dict[str, Attachment]:
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))["attachments"]
    return {key: Attachment.model_validate(value) for key, value in raw.items()}


def manifest_for(template: Path) -> Manifest:
    """The manifest beside a template: `2026-10-15.docx` → `2026-10-15.yaml`."""
    path = template.with_suffix(".yaml")
    if not path.exists():
        raise ChecklistRefused(f"{template.name}: no manifest {path.name} beside it")
    found = Manifest.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")))
    unknown = [a.key for a in found.attachments if a.key not in catalogue()]
    if unknown:
        raise ChecklistRefused(f"{path.name}: not in config/document_attachments.yaml: {unknown}")
    return found


@dataclass
class Item:
    group: Group
    key: str
    label_mk: str
    reason: Literal["applicant", "missing", "attachment"]
    documents: list[str] = field(default_factory=list)
    how_mk: str | None = None
    notes_mk: list[str] = field(default_factory=list)

    def as_row(self) -> dict:
        return {
            "group": self.group,
            "key": self.key,
            "label_mk": self.label_mk,
            "reason": self.reason,
            "documents": self.documents,
            "how_mk": self.how_mk,
            "notes_mk": self.notes_mk,
        }


def _group(key: str) -> Group:
    # From the registry: every money field is a figure; the rest by the key's own name.
    if engine.registry()[key].type == "money":
        return "money"
    if key.startswith("declaration_"):
        return "declaration"
    if key.startswith("signature_"):
        return "signature"
    return "data"


def build(documents: list[tuple[Manifest, engine.Rendered]]) -> list[Item]:
    """The package's checklist, in the order an applicant works through it."""
    items: dict[str, Item] = {}
    for manifest, rendered in documents:
        for unfilled in rendered.unfilled:
            item = items.setdefault(
                unfilled.key,
                Item(_group(unfilled.key), unfilled.key, unfilled.label_mk, unfilled.reason),
            )
            if manifest.title_mk not in item.documents:
                item.documents.append(manifest.title_mk)
        if manifest.fill_by_hand_mk:
            key = f"by-hand:{manifest.title_mk}"
            items.setdefault(
                key,
                Item("data", key, manifest.fill_by_hand_mk, "applicant", [manifest.title_mk]),
            )
        for asked in manifest.attachments:
            known = catalogue()[asked.key]
            item = items.setdefault(
                f"attachment:{asked.key}",
                Item("attachment", asked.key, known.label_mk, "attachment", how_mk=known.how_mk),
            )
            if manifest.title_mk not in item.documents:
                item.documents.append(manifest.title_mk)
            if asked.note_mk and asked.note_mk not in item.notes_mk:
                item.notes_mk.append(asked.note_mk)
    order = list(HEADINGS)
    return sorted(items.values(), key=lambda i: order.index(i.group))


def package_notes(manifests: list[Manifest]) -> list[str]:
    """Each manifest's package note, once, in order."""
    return list(dict.fromkeys(m.package_note_mk for m in manifests if m.package_note_mk))


def as_rows(items: list[Item]) -> list[dict]:
    """For `document_package.checklist`."""
    return [item.as_row() for item in items]


def to_docx(
    items: list[Item], *, applicant: str, call: str, on: dt.date, notes: list[str] = ()
) -> bytes:
    """The checklist as the package's last document. Raises ChecklistRefused on a banned phrase."""
    document = docx.Document()
    document.core_properties.language = "mk-MK"
    style = document.styles["Normal"]
    style.font.size = Pt(11)
    document.add_heading("Што уште треба да подготвите", level=1)
    document.add_paragraph(f"За: {applicant}\nЈавен повик: {call}\nСоставено на: {mkdate(on)}")
    document.add_paragraph(
        "Документите во пакетот се нацрт. Полињата означени со жолто ги пополнувате вие, "
        "со вистинските податоци на субјектот. Пред да поднесете, проверете ги барањата "
        "во последната верзија на повикот."
    )
    for note in notes:
        document.add_paragraph(note)
    for group, heading in HEADINGS.items():
        mine = [i for i in items if i.group == group]
        if not mine:
            continue
        document.add_heading(heading, level=2)
        for item in mine:
            document.add_paragraph(f"{BOX} {item.label_mk}")
            details = []
            if item.reason == "missing":
                details.append("Овој податок не го имаме; внесете го.")
            if item.how_mk:
                details.append(item.how_mk)
            details += item.notes_mk
            where = "Се бара во" if item.reason == "attachment" else "Во"
            details.append(f"{where}: {', '.join(item.documents)}.")
            detail = document.add_paragraph(" ".join(details))
            detail.paragraph_format.left_indent = Pt(18)
            detail.paragraph_format.space_after = Pt(6)

    out = io.BytesIO()
    document.save(out)
    text = engine.text_of(out.getvalue())
    banned = find_banned(text)
    if banned:
        raise ChecklistRefused(f"banned phrase in the checklist: {banned}")
    return out.getvalue()
