"""Making a template from an institution's own Word form (roadmap P5 s58).

A template is the official form, unchanged, with fields placed into it: the form is never
retyped. `place` finds a table row by the start of its Macedonian label and writes `{{r key }}`
into the row's answer cell, **only if that cell is empty**: a cell that holds anything (the
form's options, «Пополнува Агенцијата») is the institution's and is never overwritten. A label
that is not found is an error, so a changed form is noticed rather than silently half-filled.

`fill_line_above` writes the field over the blank line above a caption such as «(име и
презиме)». `append` adds the field as a new line inside a one-cell row, after its label, for forms
laid out that way. `mark_after` adds a marker paragraph after the paragraph that starts with
given text (a signature line), for forms that have no cell for it.
"""

import copy
import io
import re

import docx
from docx.document import Document
from docx.table import _Cell


class FormChanged(ValueError):
    """The form no longer has a row or line the template expects: look at the new form."""


def _label(text: str) -> str:
    # «Назив на работодавачот / Emri i punëdhënësit» → the Macedonian half; a list dash or
    # a row number before it («- Полн назив», «1. Полн назив») is not part of the label.
    return text.split(" / ")[0].split("/")[0].strip().lstrip("-–•0123456789. ").lower()


def _distinct(cells: list[_Cell]) -> list[_Cell]:
    # A merged cell appears once per grid column it spans; keep each cell once.
    seen, kept = set(), []
    for cell in cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            kept.append(cell)
    return kept


def _rows(document: Document):
    for table in document.tables:
        for row in table.rows:
            yield _distinct(row.cells)
            for cell in _distinct(row.cells):
                for inner in cell.tables:
                    for inner_row in inner.rows:
                        yield _distinct(inner_row.cells)


def place(document: Document, label: str, key: str) -> None:
    """Write `{{r key }}` into the empty answer cell of the first row labelled `label`."""
    wanted = label.lower()
    for cells in _rows(document):
        labelled = [i for i, c in enumerate(cells) if _label(c.text).startswith(wanted)]
        if not labelled:
            continue
        answers = [c for c in cells[labelled[0] + 1 :] if not c.text.strip()]
        if not answers:
            raise FormChanged(f"«{label}»: the row has no empty answer cell")
        answers[0].paragraphs[0].text = f"{{{{r {key} }}}}"
        return
    raise FormChanged(f"«{label}»: no such row in the form")


def append(document: Document, label: str, key: str) -> None:
    """Add `{{r key }}` as a new line at the end of the cell labelled `label`.

    For forms whose rows are one wide cell, label first and the answer written after it:
    the label's own text is kept, and the field follows it.
    """
    wanted = label.lower()
    for cells in _rows(document):
        for cell in cells:
            if _label(cell.text).startswith(wanted):
                cell.add_paragraph(f"{{{{r {key} }}}}")
                return
    raise FormChanged(f"«{label}»: no such cell in the form")


_BLANK = re.compile(r"_{5,}")


def fill_line_above(document: Document, caption: str, key: str) -> None:
    """Write `{{r key }}` on the blank line just above `caption`.

    Declarations print their answers as a line with a caption below, «(име и презиме)». The
    line is either typed (`_____`, only that blank is replaced; the paragraph's other words and
    runs are kept) or drawn as the bottom border of an empty paragraph (the field is written
    into it).
    """
    paragraphs = document.paragraphs
    for i, below in enumerate(paragraphs):
        if not below.text.strip().startswith(caption) or i == 0:
            continue
        above = paragraphs[i - 1]
        if _BLANK.search(above.text):
            _replace_blank(above, f"{{{{r {key} }}}}")
            return
        if (
            not above.text.strip()
            and above._p.pPr is not None
            and above._p.pPr.find(docx.oxml.ns.qn("w:pBdr")) is not None
        ):
            above.add_run(f"{{{{r {key} }}}}")
            return
        # Typed blanks may sit one paragraph higher, with an empty spacer between.
        typed = [p for p in paragraphs[:i] if p.text.strip()]
        if typed and _BLANK.search(typed[-1].text):
            _replace_blank(typed[-1], f"{{{{r {key} }}}}")
            return
    raise FormChanged(f"«{caption}»: no blank line above such a caption in the form")


def _replace_blank(paragraph, text: str) -> None:
    # The blank may be split across runs; cut it out character by character, run by run,
    # and put the text where it began.
    start, end = _BLANK.search(paragraph.text).span()
    position, placed = 0, False
    for run in paragraph.runs:
        value = run.text
        run_start, run_end = position, position + len(value)
        position = run_end
        if run_end <= start or run_start >= end:
            continue
        cut_from, cut_to = max(start, run_start) - run_start, min(end, run_end) - run_start
        insert = "" if placed else text
        run.text = value[:cut_from] + insert + value[cut_to:]
        placed = True


def mark_after(document: Document, starts_with: str, key: str) -> None:
    """Add a paragraph holding `{{r key }}` after the first paragraph starting `starts_with`."""
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(starts_with):
            new = copy.deepcopy(paragraph._p)
            paragraph._p.addnext(new)
            for run in new.findall(docx.oxml.ns.qn("w:r"))[1:]:
                new.remove(run)
            added = docx.text.paragraph.Paragraph(new, paragraph._parent)
            added.text = f"{{{{r {key} }}}}"
            return
    raise FormChanged(f"«{starts_with}»: no such line in the form")


def to_bytes(document: Document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
