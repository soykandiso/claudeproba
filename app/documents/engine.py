"""Filling a .docx template, never with an invented value (roadmap P5 s56, brief §8).

A template is a Word file whose fields are written `{{r key }}`, every key one of
`config/document_fields.yaml`. The registry says where each field's value may come from;
this module makes that a rule rather than a habit:

- a field whose source is **`applicant`** (money, exact counts, people, declarations,
  signatures) is **never filled**. It prints «[Пополнува барателот: <label>]», highlighted.
  Handing it a value is a bug, and is refused (`InventedValue`), not ignored;
- any other field with no value prints «[Недостасува: <label>]», highlighted, and is listed;
- a template naming a field the registry does not know, or writing a field without the `r`
  (docxtpl then drops a highlighted marker without a trace: found when this was built), is
  refused before anything is filled (`TemplateRefused`).

`render` returns the document and both lists, `filled` and `unfilled`, so a package can
carry what the applicant must still supply (`generated_document.unfilled_fields`, s57).

The output is a draft for human completion (brief §8). Saying so on the document itself is
s66; a .pdf of it is later, with the package (s67).
"""

import io
import re
from dataclasses import dataclass
from decimal import Decimal
from functools import cache
from pathlib import Path
from typing import Literal

import docx
import yaml
from docxtpl import DocxTemplate, RichText
from pydantic import BaseModel, ConfigDict

from app.i18n import number
from app.web.format import mkdate

REGISTRY = Path(__file__).resolve().parents[2] / "config" / "document_fields.yaml"
MARK = "FFF2A8"  # the marker's shading: a pale yellow that prints as grey, readable either way

APPLICANT_SAYS = "Пополнува барателот"
MISSING_SAYS = "Недостасува"

# `{{r key }}` or `{{ key }}`, as written in the paragraph's visible text.
_PLACEHOLDER = re.compile(r"\{\{\s*(r\s+)?([^\s}|]+)[^}]*\}\}")
_CONTROL = re.compile(r"\{%")


class Field(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    label_mk: str
    source: Literal["buyer", "profile", "call", "system", "applicant"]
    type: Literal["text", "date", "money", "number"]


class TemplateRefused(ValueError):
    """The template cannot be filled safely; nothing was produced."""


class InventedValue(ValueError):
    """A value was handed for a field only the applicant may supply."""


@dataclass(frozen=True)
class Unfilled:
    key: str
    label_mk: str
    reason: Literal["applicant", "missing"]

    def as_row(self) -> dict:
        return {"key": self.key, "label_mk": self.label_mk, "reason": self.reason}


@dataclass(frozen=True)
class Rendered:
    docx: bytes
    filled: list[str]
    unfilled: list[Unfilled]


@cache
def registry(path: Path = REGISTRY) -> dict[str, Field]:
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))["fields"]
    return {key: Field.model_validate(value) for key, value in raw.items()}


def _texts(document: docx.document.Document) -> list[str]:
    """The visible text of every paragraph a template can hold a field in."""
    parts = [document.element.body]
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if not part.is_linked_to_previous:
                parts.append(part._element)
    texts = []
    for element in parts:
        for paragraph in element.iter(docx.oxml.ns.qn("w:p")):
            texts.append("".join(t.text or "" for t in paragraph.iter(docx.oxml.ns.qn("w:t"))))
    return texts


def placeholders(template: bytes) -> list[tuple[str, bool]]:
    """Every field a template names, in order, and whether it is written with the `r`."""
    found = []
    for text in _texts(docx.Document(io.BytesIO(template))):
        found += [(key, bool(r)) for r, key in _PLACEHOLDER.findall(text)]
    return found


def check(template: bytes) -> list[str]:
    """Why the template cannot be filled, one line each; empty when it can."""
    problems, fields = [], registry()
    document = docx.Document(io.BytesIO(template))
    if any(_CONTROL.search(text) for text in _texts(document)):
        problems.append("a {% … %} block: a template here holds fields only, no logic")
    for key, rich in placeholders(template):
        if key not in fields:
            problems.append(f"{{{{r {key} }}}}: not a field in config/document_fields.yaml")
        elif not rich:
            problems.append(f"{{{{ {key} }}}}: write it {{{{r {key} }}}}, or a marker is lost")
    return list(dict.fromkeys(problems))


def _marker(says: str, label: str) -> RichText:
    return RichText(f"[{says}: {label}]", bold=True, highlight=MARK)


def _formatted(field: Field, value) -> str:
    if field.type == "date":
        return mkdate(value)
    if field.type == "money":
        return f"{number(Decimal(str(value)), 2)} МКД"
    if field.type == "number":
        return number(value)
    return str(value)


def _empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def render(template: bytes, values: dict) -> Rendered:
    """Fill `template` from `values`; refuse rather than invent. See the module's rules."""
    fields = registry()
    problems = check(template)
    if problems:
        raise TemplateRefused("; ".join(problems))
    unknown = sorted(set(values) - set(fields))
    if unknown:
        raise TemplateRefused(f"values for fields that do not exist: {', '.join(unknown)}")
    invented = sorted(
        k for k, v in values.items() if fields[k].source == "applicant" and not _empty(v)
    )
    if invented:
        raise InventedValue(f"only the applicant can supply: {', '.join(invented)}")

    context, filled, unfilled = {}, [], []
    for key in dict.fromkeys(k for k, _ in placeholders(template)):
        field, value = fields[key], values.get(key)
        if field.source == "applicant":
            context[key] = _marker(APPLICANT_SAYS, field.label_mk)
            unfilled.append(Unfilled(key, field.label_mk, "applicant"))
        elif _empty(value):
            context[key] = _marker(MISSING_SAYS, field.label_mk)
            unfilled.append(Unfilled(key, field.label_mk, "missing"))
        else:
            context[key] = RichText(_formatted(field, value))
            filled.append(key)

    document = DocxTemplate(io.BytesIO(template))
    document.render(context, autoescape=True)
    out = io.BytesIO()
    document.save(out)
    return Rendered(out.getvalue(), filled, unfilled)


def text_of(document: bytes) -> str:
    """Every paragraph's text, one per line: what a reader of the file sees, for checks."""
    return "\n".join(_texts(docx.Document(io.BytesIO(document))))
